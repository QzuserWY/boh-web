"""Read-only, sheet-specific import. Run from any directory; never changes the workbook."""
import json, re
from pathlib import Path
from collections import Counter
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
w = openpyxl.load_workbook(ROOT.parent / '司辰之书.xlsx', data_only=True)
planner = json.loads((ROOT / 'src/data.json').read_text(encoding='utf-8'))
ASPECTS = list('刃铸杯心启灯月蛾蜜引鳞穹冬')
entries = []
def val(sheet, r, c):
    x = w[sheet].cell(r,c).value
    return '' if x is None else str(x).strip()
def nums(sheet,r,start=4):
    return {a:float(val(sheet,r,start+i)) for i,a in enumerate(ASPECTS) if re.fullmatch(r'\d+(\.\d+)?',val(sheet,r,start+i)) and float(val(sheet,r,start+i))>0}
def parsed(s):
    return {a:int(n) if n else None for a,n in re.findall(r'([刃铸杯心启灯月蛾蜜引鳞穹冬])(\d*)',s)}
def add(cat,name,sheet,loc,fields=None,aspects=None,branches=None,**kw):
    e=dict(id=f'{sheet}:{loc}:{name}',category=cat,name=name,fields=[dict(label=k,value=v) for k,v in (fields or {}).items() if v],aspects=aspects or {},branches=branches or [],sources=[f'{sheet}!{loc}'],sections=[],spoiler=False,**kw)
    entries.append(e)
    return e
def table(e,title,cols,rows):
    e['sections'].append(dict(title=title,columns=cols,rows=rows))
skill_entries={}
for s in planner['skills']:
    e=add('技能',s['name'],'上树',s['sourceCell'],{'别名':'、'.join(s['aliases']),'原始规划':f"{s['original']['branch']} · {s['original']['rank']} 档",'分支与魂质':'；'.join(o['branch']+' → '+o['soul'] for o in s['options'])},branches=[o['branch'] for o in s['options']],skillId=s['id'])
    skill_entries[s['id']]=e
    skill_entries[s['name']]=e

# Recipe columns form independent vertical blocks. Unpaired requirements remain visible as incomplete records.
for c in range(1,15,2):
    owner=None
    for r in range(2,55):
        name,req=val('技艺',r,c),val('技艺',r,c+1)
        if name and re.fullmatch('[刃铸杯心启灯月蛾蜜引鳞穹冬]、[刃铸杯心启灯月蛾蜜引鳞穹冬]',req):
            owner=name
            if name in skill_entries:
                skill_entries[name]['aspects']=parsed(req)
                skill_entries[name]['sources'].append(f'技艺!{w["技艺"].cell(r,c).coordinate}:{w["技艺"].cell(r,c+1).coordinate}')
        elif owner and re.match(r'[刃铸杯心启灯月蛾蜜引鳞穹冬](5|10|15)(?:\s|$)',req):
            loc=f'{w["技艺"].cell(r,c).coordinate}:{w["技艺"].cell(r,c+1).coordinate}'
            e=add('制作配方',name or '原表未填写产物','技艺',loc,{'使用技能':owner,'制作需求':req,'说明':'5 / 10 / 15 是准则强度要求，不是技能等级。'},parsed(re.match(r'^.\d+',req).group()),skillId=skill_entries.get(owner,{}).get('skillId',''))
            if not name or name in req: e['fields'].append(dict(label='原表疑点',value='产物缺失或与材料同名，保留原文，请以游戏和 Wiki 为准。'))

for r in list(range(3,26))+list(range(27,38))+list(range(39,54))+list(range(55,68))+list(range(69,78)):
    name=val('回忆',r,1)
    if not name: continue
    br=[b for b in planner['branches'] if b in val('回忆',r,17)]
    e=add('天气' if r>=69 else '回忆',name,'回忆',f'A{r}:Q{r}',{'获取／原表提示':val('回忆',r,3),'用于进阶':'、'.join(br),'类型':'闰识' if 55<=r<=67 else ('天气' if r>=69 else '回忆')},nums('回忆',r),br)
    e['spoiler']=55<=r<=67
for r in range(2,81):
    name=val('工作台',r,1)
    if not name: continue
    raw=val('工作台',r,3).replace('鸟鸣术','鸟鸣学')
    add('工作台',name,'工作台',f'A{r}:R{r}',{'类型':val('工作台',r,2),'进阶分支':raw,'第四卡槽':val('工作台',r,17),'第五卡槽':val('工作台',r,18),'使用提示':'分支相符仍需检查准则和卡槽。除床以外，工作台拒绝疲劳；书桌拒绝胶卷和唱片。'},parsed(''.join(val('工作台',r,c) for c in range(4,17))),[b for b in planner['branches'] if b in raw])
for r in [2,3,6,7,8,9,10,13,14,*range(18,26)]:
    add('物品工具',val('工具',r,1),'工具',f'A{r}:F{r}',{'位置／说明':val('工具',r,6)}, {a:n or 1 for a,n in parsed(''.join(val('工具',r,c) for c in range(2,6))).items()})
for r in range(2,16):
    add('物品工具',val('工具',r,8),'工具',f'H{r}:U{r}',{'分组':'工具页 · 墨水与物品'},nums('工具',r,9))
food={}
for sheet in ['饮食','食谱']:
    for r in range(2,168):
        name=val(sheet,r,1)
        if not name or not val(sheet,r,2): continue
        fields={'定位':val(sheet,r,2),'类别':val(sheet,r,3),'原料／获得方式 1':val(sheet,r,17),'原料／获得方式 2':val(sheet,r,18),'工具':val(sheet,r,19),'菜谱来源' if sheet=='饮食' else '额外信息':val(sheet,r,20)}
        if sheet=='饮食': fields['额外信息']=val(sheet,r,21)
        if name in food:
            e=food[name];e['sources'].append(f'{sheet}!A{r}:T{r}')
            table(e,'食谱页对照（保留原表差异）',['项目','原文'],[[k,v] for k,v in fields.items() if v]+[['准则',' '.join(a+str(n) for a,n in nums(sheet,r).items())]])
        else: food[name]=add('饮食',name,sheet,f'A{r}:{"U" if sheet=="饮食" else "T"}{r}',fields,nums(sheet,r))
for r in range(17,31):
    add('助手',val('助手',r,2),'助手',f'B{r}:F{r}',{'特殊增益':val('助手',r,4),'邀请条件':val('助手',r,5),'魂质参考':val('助手',r,6)},parsed(val('助手',r,3)))
for r in range(2,15):
    add('助手增益',val('助手',r,1)+' · 增益参考','助手',f'A{r}:H{r}',{val('助手',1,c) or ['','助手','魂质','存留回忆','工具','饮品','易获取食物','特殊增益'][c-1]:val('助手',r,c) for c in range(2,9)}, {val('助手',r,1):None})
for r in range(2,25):
    e=add('访客',val('访客',r,1),'访客',f'A{r}:AC{r}',{'邀请所需':val('访客',r,2),'饮食偏好':val('访客',r,3),'额外兴趣':val('访客',r,4),'谈话魂质':'、'.join(filter(None,[val('访客',r,5),val('访客',r,6)]))})
    table(e,'沙龙谈话收获',['搭配访客（原表简称）','技能／次数'],[[val('访客',1,c),val('访客',r,c)] for c in range(7,30) if val('访客',r,c) not in ('','\\','/')])
def grid(name,start,end,c1,c2,header=None,note=''):
    e=add('探索',name,'探索',f'{w["探索"].cell(start,c1).coordinate}:{w["探索"].cell(end,c2).coordinate}',{'说明':note})
    table(e,name,header or [val('探索',start,c) or '原表空列' for c in range(c1,c2+1)],[[val('探索',r,c) for c in range(c1,c2+1)] for r in range(start+1,end+1) if any(val('探索',r,c) for c in range(c1,c2+1))])
for name,start,end,c1,c2 in [('花园采集',6,12,1,6),('拾滩',14,26,1,6),('荒野',29,38,1,6),('蜂巢',40,45,1,6),('海鸥巢群',49,58,2,6),('海蚀洞',61,67,2,6)]:
    grid(name,start,end,c1,c2,note='花园原表物品和准则混排，按原样展示，不推断季节归属。' if start==6 else '')
for c,start,end in [(8,14,17),(9,14,22),(10,14,18),(11,14,29),(8,21,25),(8,29,45),(9,29,39),(10,29,37)]:grid(val('探索',start,c),start,end,c,c,['可能获得（原表）'])
grid('凯特与赫洛有限公司订单',49,66,8,9,['商品','价格（便士）'])
grid('T.R.N.有限公司订单',49,63,10,12,['商品','价格（便士）','备注'])
for r in range(1,4):add('探索',val('探索',r,1).split('：')[0],'探索',f'A{r}',{'原表提示':val('探索',r,1)})
for name,asp in {'明识':'灯2穹1','刚毅':'铸2刃1','灵躯':'引2蛾1','辩闻':'启2铸1','黯晦':'杯2刃1','脉律':'心2杯1','悲恸':'蛾2月1','铭晓':'冬2灯1','健康':'心1蜜1鳞1'}.items():
    add('魂质',name,'助手','A2:C14',{'说明':'基础魂质；技能上树的魂质收益可在技能条目查看。'},parsed(asp))
guides=[('上树与调谐','每个技能有两条可选智慧分支，只能呈递一次。上树位置需要满足等级与槽位条件。本工具将未来计划和实际上树分开记录。已调谐表示已用该技能完成一次魂质进阶。','新手指南'),('魂质进阶','需要两个同种、同等级且未疲劳的魂质，以及已经呈递、收益匹配且未用过的技能。工作台必须接受投入的准则和物品；相应进阶条件可由工作台或回忆提供。仅分支相同不能保证可用。','魂质'),('教诲与技能升级','升级需要匹配技能准则的教诲，也可搭配匹配准则的其他回忆。百科中的制作强度 5、10、15 不是技能等级。','FAQ'),('制作与回忆','配方按准则强度分为 5、10、15 等门槛，高阶配方还可能要求指定材料。天气也能作为回忆使用；普通非存留回忆会在黎明消逝。','新手指南')]
for name,body,page in guides:
    e=add('机制指南',name,'Wiki',page,{'规则摘要':body});e['wiki']='https://boh.huijiwiki.com/wiki/'+page
out=dict(entries=entries,counts=dict(Counter(e['category'] for e in entries)),source='司辰之书.xlsx',checked='2026-09-22')
(ROOT/'src/encyclopedia-data.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(out['counts']);print('Total',len(entries))
