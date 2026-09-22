"""Read the user's original workbook without modifying it; build the bundled plan."""
import openpyxl, json, re
from pathlib import Path

root = Path(__file__).resolve().parents[1]
w = openpyxl.load_workbook(root.parent / '司辰之书.xlsx')
s = w['上树']
groups = {
 '司辰学': {'刚毅':'鼓点与舞步 洪钟与铜器 力量之墨 利米亚典仪 收容之墨 无侵之敕令 武力之敕令 诱惑与消殒 渊深曼达安语', '灵躯':'雕珀与琢石 窥天术 门扉与墙垣 萨巴佐因语 天空的故事 伟大符印与伟大伤疤 无敌太阳典仪'},
 '盗火术': {'明识':'玻璃吹制与容器制造 玻璃与锻光 雕珀与琢石 精炼与擢升 启示之墨 嬗变与解放 渊深曼达安语', '铭晓':'缝合与装订 腐化与煅烧 富奇诺语 赫沃尔语 洪钟与铜器 阙前格律 溶解与离析 锁簧与发条 香料与滋味 焰篆体'},
 '照明术': {'刚毅':'玻璃与锻光 锤砧的戒律 光明果悦音 精炼与擢升 嬗变与解放 焰篆体', '明识':'残破十字路 厄里卡帕奥语 弧月的倒影 解踪语 阙前格律 伤疤的戒律 守夜人的悖论 曙光的静观 无敌太阳典仪 武力之敕令'},
 '静默术': {'悲恸':'伤疤的戒律 富奇诺语 力量之墨 疫疠之秘 窥天术 淬火与熄焰 利米亚典仪 群蛇与毒液 伟大符印与伟大伤疤 厄里卡帕奥语', '铭晓':'玻璃吹制与容器制造 启示之墨 锋锐 景象与感知 白雪的故事 手术与放血'},
 '夜游术': {'辩闻':'锤砧的戒律 大海的故事 镰刀与日蚀 门扉与墙垣 虬蟠与裂隙 希克索斯语 弦乐与歌谣', '灵躯':'道途与朝圣 伐诃语 景象与感知 昆虫与花蜜 毛皮与羽翎 守夜人的悖论 曙光的静观 阈限之敕令 珍珠与浪潮'},
 '蠕虫学': {'辩闻':'残破十字路 锋锐 群蛇与毒液 兰花与迷药 溶解与离析 沙的故事 兽角与象牙 锁簧与发条', '黯晦':'覆画残迹与古老前身 腐化与煅烧 狼的故事 虬蟠与裂隙 希克索斯语 夜魄语 音律与铭记 阈限之敕令'},
 '丛林学': {'健康':'缝纫与编织 复苏与羽化 鼓点与舞步 赫沃尔语 毛皮与羽翎 香料与滋味 叶片与棘刺 夜魄语 珍珠与浪潮', '黯晦':'草药与药汤 根系之仪 光明果悦音 昆虫与花蜜 兰花与迷药 三重誓之仪 兽角与象牙 诱惑与消殒'},
 '保存术': {'脉律':'解踪语 拉姆桑德语 镰刀与日蚀 手术与放血 缝合与装订 复苏与羽化 岩石的故事 无侵之敕令', '健康':'收容之墨 疫疠之秘 草药与药汤 三重誓之仪 丘陵与孔窍 弧月的倒影 淬火与熄焰 根系之仪'},
 '鸟鸣学': {'脉律':'缝纫与编织 叶片与棘刺 丘陵与孔窍 大海的故事 道途与朝圣 伐诃语 覆画残迹与古老前身 弦乐与歌谣', '悲恸':'拉姆桑德语 岩石的故事 白雪的故事 狼的故事 萨巴佐因语 沙的故事 天空的故事 音律与铭记'},
}
aliases = {'虬蟠与裂缝':'虬蟠与裂隙', '疫疠的奥秘':'疫疠之秘'}
rules = {}
for branch, souls in groups.items():
 for soul, names in souls.items():
  for name in names.split(): rules.setdefault(name, []).append({'branch':branch,'soul':soul})
colors = {'刚毅':'#FFC000','明识':'#FFFF00','铭晓':'#C78FFF','悲恸':'#8FAADC','灵躯':'#00B0F0','黯晦':'#7030A0','辩闻':'#D60093','健康':'#FF6D6D','脉律':'#FF99CC'}
skills=[]
for row in range(2,11):
 for col in range(2,11):
  c=s.cell(row,col)
  if not c.value: continue
  name=re.sub(r'\d+$','',str(c.value).strip())
  canonical=aliases.get(name,name)
  options=rules[canonical]
  assert len(options)==2,(canonical,options)
  branch=s.cell(1,col).value
  soul=next(o['soul'] for o in options if o['branch']==branch)
  color=c.font.color
  red=bool(color and color.type=='rgb' and color.rgb in ['FFFF0000','00FF0000'])
  skills.append({'id':canonical,'name':name,'aliases':[canonical] if canonical!=name else [],'options':options,'original':{'branch':branch,'rank':row-1},'sourceCell':c.coordinate,'sourceSoul':soul,'initialAttuned':red})
assert len(skills)==73 and len({x['id'] for x in skills})==73
assert set(rules)=={x['id'] for x in skills}, set(rules)-{x['id'] for x in skills}
out={'branches':list(groups),'souls':colors,'skills':skills,'source':{'workbook':'司辰之书.xlsx','range':'上树!B2:J10','wiki':'https://boh.huijiwiki.com/wiki/列表/魂质获得','verifiedAt':'2026-09-22'}}
(root/'src').mkdir(exist_ok=True)
(root/'src'/'data.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Imported {len(skills)} skills, {sum(len(x["options"]) for x in skills)} verified branch options. Original workbook unchanged.')
